"""KDP royalty / order reports (CSV, or XLSX when openpyxl is installed) -> per-title totals.

Download from KDP Reports: Royalties Estimator or Orders -> 'Download report'.
"""

from __future__ import annotations

import csv
import datetime as dt
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

from ..niche.models import parse_date, parse_float, parse_int

ALIASES: dict[str, list[str]] = {
    "date": ["royalty date", "date", "order date", "transaction date", "month"],
    "title": ["title", "book title"],
    "author": ["author name", "author"],
    "asin": ["asin/isbn", "asin", "isbn", "asin / isbn"],
    "marketplace": ["marketplace", "market", "store"],
    "royalty_type": ["royalty type", "type"],
    "transaction_type": ["transaction type"],
    "units": ["net units sold", "units sold", "units", "net units", "orders"],
    "refunds": ["units refunded", "refunds", "returns"],
    "list_price": ["avg. list price without tax", "list price", "avg list price"],
    "royalty": ["royalty", "royalty (usd)", "earnings", "royalty amount"],
    "currency": ["currency"],
    "kenp": ["kenp read", "kenp", "pages read", "kenp pages read"],
}


def _norm(h: str) -> str:
    return re.sub(r"\s+", " ", str(h).strip().lower().replace("_", " "))


def column_map(headers: list[str]) -> dict[str, str]:
    normalized = {_norm(h): h for h in headers}
    out: dict[str, str] = {}
    for field_name, aliases in ALIASES.items():
        for a in aliases:
            if a in normalized and field_name not in out:
                out[field_name] = normalized[a]
                break
    return out


@dataclass
class SaleRow:
    date: dt.date | None
    title: str
    asin: str | None
    marketplace: str | None
    royalty_type: str | None
    units: int
    refunds: int
    royalty: float
    currency: str | None
    kenp: int = 0

    @property
    def month(self) -> str:
        return self.date.strftime("%Y-%m") if self.date else "unknown"


def _rows_from_file(path: Path) -> list[dict]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Reading .xlsx needs openpyxl: pip install 'kdp-studio[xlsx]' (or export the report as CSV)") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        best: list[dict] = []
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            # find the header row: the first row that maps to a title + royalty/units column
            for i, row in enumerate(rows[:20]):
                headers = [str(c) if c is not None else "" for c in row]
                cmap = column_map(headers)
                if "title" in cmap and ("royalty" in cmap or "units" in cmap):
                    data = [dict(zip(headers, r)) for r in rows[i + 1 :] if any(c is not None for c in r)]
                    if len(data) > len(best):
                        best = data
                    break
        return best
    text = path.read_text(encoding="utf-8-sig")
    return list(csv.DictReader(text.splitlines()))


def load_sales(path: Path) -> list[SaleRow]:
    rows = _rows_from_file(path)
    if not rows:
        return []
    cmap = column_map(list(rows[0].keys()))
    if "title" not in cmap or ("royalty" not in cmap and "units" not in cmap):
        raise ValueError(f"Could not find title and royalty/units columns. Headers: {list(rows[0].keys())}")
    out: list[SaleRow] = []
    for r in rows:
        g = lambda f: r.get(cmap[f]) if f in cmap else None  # noqa: E731
        title = (str(g("title") or "")).strip()
        if not title or title.lower().startswith("total"):
            continue
        out.append(
            SaleRow(
                date=parse_date(g("date")),
                title=title,
                asin=(str(g("asin")).strip() if g("asin") else None),
                marketplace=(str(g("marketplace")).strip() if g("marketplace") else None),
                royalty_type=(str(g("royalty_type")).strip() if g("royalty_type") else None),
                units=parse_int(g("units")) or 0,
                refunds=parse_int(g("refunds")) or 0,
                royalty=parse_float(g("royalty")) or 0.0,
                currency=(str(g("currency")).strip() if g("currency") else None),
                kenp=parse_int(g("kenp")) or 0,
            )
        )
    return out


def summarize_sales(rows: list[SaleRow], since: dt.date | None = None) -> dict:
    if since:
        rows = [r for r in rows if r.date is None or r.date >= since]
    by_title: dict[str, dict] = defaultdict(lambda: {"units": 0, "refunds": 0, "royalty": 0.0, "kenp": 0, "marketplaces": defaultdict(float), "months": defaultdict(float), "asins": set()})
    by_month: dict[str, dict] = defaultdict(lambda: {"units": 0, "royalty": 0.0, "kenp": 0})
    for r in rows:
        t = by_title[r.title]
        t["units"] += r.units
        t["refunds"] += r.refunds
        t["royalty"] += r.royalty
        t["kenp"] += r.kenp
        t["marketplaces"][r.marketplace or "?"] += r.royalty
        t["months"][r.month] += r.royalty
        if r.asin:
            t["asins"].add(r.asin)
        m = by_month[r.month]
        m["units"] += r.units
        m["royalty"] += r.royalty
        m["kenp"] += r.kenp
    titles = [
        {"title": k, "units": v["units"], "refunds": v["refunds"], "royalty": round(v["royalty"], 2), "kenp": v["kenp"], "asins": sorted(v["asins"]), "marketplaces": dict(v["marketplaces"]), "months": dict(v["months"])}
        for k, v in by_title.items()
    ]
    titles.sort(key=lambda t: -t["royalty"])
    return {
        "rows": len(rows),
        "since": since.isoformat() if since else None,
        "units": sum(r.units for r in rows),
        "refunds": sum(r.refunds for r in rows),
        "royalty": round(sum(r.royalty for r in rows), 2),
        "kenp": sum(r.kenp for r in rows),
        "titles": titles,
        "months": dict(sorted(by_month.items())),
    }


def sales_markdown(summary: dict) -> str:
    lines = ["## KDP sales", ""]
    lines.append(f"Rows: {summary['rows']} · net units {summary['units']:,} (refunds {summary['refunds']}) · royalty {summary['royalty']:,.2f} · KENP {summary['kenp']:,}" + (f" · since {summary['since']}" if summary["since"] else ""))
    lines.append("")
    lines.append("| Title | ASIN/ISBN | Units | Refunds | Royalty | KENP | Top marketplace |")
    lines.append("|---|---|---|---|---|---|---|")
    for t in summary["titles"]:
        top_mp = max(t["marketplaces"].items(), key=lambda kv: kv[1])[0] if t["marketplaces"] else "–"
        lines.append(f"| {t['title']} | {', '.join(t['asins']) or '–'} | {t['units']:,} | {t['refunds']} | {t['royalty']:,.2f} | {t['kenp']:,} | {top_mp} |")
    if summary["months"]:
        lines.append("")
        lines.append("| Month | Units | Royalty | KENP |")
        lines.append("|---|---|---|---|")
        for m, v in summary["months"].items():
            lines.append(f"| {m} | {v['units']:,} | {v['royalty']:,.2f} | {v['kenp']:,} |")
    return "\n".join(lines) + "\n"
